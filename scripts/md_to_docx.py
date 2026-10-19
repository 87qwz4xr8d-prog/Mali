#!/usr/bin/env python3
"""Convert simple markdown ISO forms to .docx for Word/LibreOffice."""

from __future__ import annotations

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


CHECKBOX = "☐"
CHECKED = "☑"


def set_run_font(run, *, bold=False, size=11, name="TH Sarabun New"):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:eastAsia"), name)


def add_formatted_paragraph(doc, text: str, *, style=None, bold_all=False, size=11, space_after=6):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)

    if bold_all:
        run = p.add_run(text)
        set_run_font(run, bold=True, size=size)
        return p

    parts = re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = p.add_run(part[2:-2])
            set_run_font(run, bold=True, size=size)
        elif part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            set_run_font(run, size=size, name="Consolas")
        else:
            run = p.add_run(part.replace(CHECKBOX, "☐").replace(CHECKED, "☑"))
            set_run_font(run, size=size)
    return p


def parse_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_separator_row(cells: list[str]) -> bool:
    if not cells:
        return False
    return all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells if c != "")


def add_table(doc, rows: list[list[str]]):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.rows[i].cells[j]
            text = row[j] if j < len(row) else ""
            cell.text = ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            parts = re.split(r"(\*\*[^*]+\*\*)", text)
            for part in parts:
                if not part:
                    continue
                if part.startswith("**") and part.endswith("**"):
                    run = p.add_run(part[2:-2])
                    set_run_font(run, bold=True, size=10)
                else:
                    run = p.add_run(part)
                    set_run_font(run, bold=(i == 0), size=10)
            if i == 0:
                shading = cell._element.get_or_add_tcPr()
                shd = shading.makeelement(
                    qn("w:shd"),
                    {
                        qn("w:fill"): "E8EEF5",
                        qn("w:val"): "clear",
                    },
                )
                shading.append(shd)
    doc.add_paragraph()


def convert_md_to_docx(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)

    style = doc.styles["Normal"]
    style.font.name = "TH Sarabun New"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "TH Sarabun New")

    i = 0
    in_code = False
    code_lines: list[str] = []
    table_rows: list[list[str]] = []

    def flush_table():
        nonlocal table_rows
        if table_rows:
            add_table(doc, table_rows)
            table_rows = []

    def flush_code():
        nonlocal code_lines
        if not code_lines:
            return
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run("\n".join(code_lines))
        set_run_font(run, size=10, name="Consolas")
        code_lines = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            flush_table()
            if in_code:
                flush_code()
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if stripped.startswith("|"):
            cells = parse_table_row(stripped)
            if is_separator_row(cells):
                i += 1
                continue
            table_rows.append(cells)
            i += 1
            continue
        else:
            flush_table()

        if not stripped:
            i += 1
            continue

        if stripped.startswith("---"):
            i += 1
            continue

        if stripped.startswith("# "):
            p = add_formatted_paragraph(doc, stripped[2:], bold_all=True, size=18, space_after=4)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x1A, 0x3A, 0x5C)
            i += 1
            continue

        if stripped.startswith("## "):
            add_formatted_paragraph(doc, stripped[3:], bold_all=True, size=14, space_after=6)
            i += 1
            continue

        if stripped.startswith("### "):
            add_formatted_paragraph(doc, stripped[4:], bold_all=True, size=12, space_after=4)
            i += 1
            continue

        if re.match(r"^\d+\.\s", stripped):
            add_formatted_paragraph(doc, stripped, size=11, space_after=3)
            i += 1
            continue

        if stripped.startswith("- "):
            add_formatted_paragraph(doc, "• " + stripped[2:], size=11, space_after=3)
            i += 1
            continue

        add_formatted_paragraph(doc, stripped, size=11, space_after=6)
        i += 1

    flush_table()
    flush_code()

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("inputs", nargs="+", type=Path)
    parser.add_argument("--outdir", type=Path, default=None)
    args = parser.parse_args()

    for md in args.inputs:
        outdir = args.outdir or md.parent
        out = outdir / (md.stem + ".docx")
        convert_md_to_docx(md, out)
        print(f"Wrote {out}")


if __name__ == "__main__":
    main()
